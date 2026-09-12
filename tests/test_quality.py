"""quality：analyze 负几何检测 + render_with_quality 修复循环（mock）。"""
import copy

import src.quality as q
import src.renderer as rd
from src.models import DeckPlan, SlideData


def _plan():
    return DeckPlan(theme_name="business", slides=[
        SlideData(kind="title"), SlideData(kind="bullets", bullets=["a", "b"]),
        SlideData(kind="thanks")])


def test_analyze_detects_negative_geometry(tmp_path):
    from pptx import Presentation
    from pptx.oxml.ns import qn
    from pptx.util import Emu

    out = str(tmp_path / "bad.pptx")
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    sh = slide.shapes.add_shape(1, 0, 0, Emu(914400), Emu(914400))
    xfrm = sh._element.spPr.find(qn("a:xfrm"))
    ext = xfrm.find(qn("a:ext"))
    ext.set("cy", str(-914400))  # 高度改为负（非法 OOXML）
    prs.save(out)

    issues = q.analyze(out)
    assert any(i.kind == "negative_geometry" for i in issues)


def test_render_with_quality_repair_loop(monkeypatch, tmp_path):
    plan = _plan()
    state = {"renders": 0, "repairs": 0}

    def fake_render(p, path, **kw):
        state["renders"] += 1
        return path

    def fake_analyze(path):
        return [q.Issue(1, "overflow_x", 1.0, "forced")]

    def repairer(p, issues):
        state["repairs"] += 1
        return copy.deepcopy(p)

    monkeypatch.setattr(rd, "render", fake_render)
    monkeypatch.setattr(q, "analyze", fake_analyze)
    out, issues, tries = q.render_with_quality(
        plan, str(tmp_path / "x.pptx"), max_tries=2, quiet=True,
        repairer=repairer, max_repair_rounds=1)
    assert state["renders"] >= 2
    assert state["repairs"] == 1
    assert issues


def test_render_with_quality_none_repairer_no_recursion(monkeypatch, tmp_path):
    plan = _plan()
    state = {"renders": 0, "repairs": 0}

    def fake_render(p, path, **kw):
        state["renders"] += 1
        return path

    monkeypatch.setattr(rd, "render", fake_render)
    monkeypatch.setattr(q, "analyze", lambda path: [q.Issue(1, "overflow_x", 1.0, "f")])
    out, issues, tries = q.render_with_quality(
        plan, str(tmp_path / "x.pptx"), max_tries=2, quiet=True,
        repairer=lambda p, i: None)
    assert state["renders"] == 2   # 无修复器返回 -> 不递归
    assert issues


def test_render_with_quality_hard_fail_raises(monkeypatch, tmp_path):
    plan = _plan()

    monkeypatch.setattr(rd, "render", lambda p, path, **kw: path)
    monkeypatch.setattr(q, "analyze", lambda path: [q.Issue(1, "overflow_x", 1.0, "f")])
    try:
        q.render_with_quality(
            plan, str(tmp_path / "x.pptx"), max_tries=2, quiet=True,
            repairer=lambda p, i: None, hard_fail=True)
        assert False, "expected LayoutConstraintError"
    except q.LayoutConstraintError as e:
        assert e.issues


def test_render_with_quality_hard_fail_false_keeps_old_behavior(monkeypatch, tmp_path):
    plan = _plan()

    monkeypatch.setattr(rd, "render", lambda p, path, **kw: path)
    monkeypatch.setattr(q, "analyze", lambda path: [q.Issue(1, "overflow_x", 1.0, "f")])
    out, issues, tries = q.render_with_quality(
        plan, str(tmp_path / "x.pptx"), max_tries=2, quiet=True,
        repairer=lambda p, i: None, hard_fail=False)
    assert issues
