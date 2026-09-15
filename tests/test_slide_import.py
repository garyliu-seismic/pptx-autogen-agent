"""slide_import：跨 pptx 整页幻灯片克隆拼接（OPC part 级），以及配套的
remove_slide 工具函数。"""
import os

from PIL import Image
from pptx import Presentation
from pptx.util import Inches

import re
import zipfile

from src.slide_import import (
    _PartnameAllocator, fix_app_properties_slide_count, remove_slide,
    splice_external_slides,
)


def _slide_texts(slide) -> list[str]:
    return [sh.text_frame.text.strip() for sh in slide.shapes
            if sh.has_text_frame and sh.text_frame.text.strip()]


def _make_src_pptx(tmp_path, texts: list[str], with_image: bool = False) -> str:
    """造一个真实的 pptx，每个 texts[i] 对应一页（同一个 layout）。"""
    prs = Presentation()
    layout = prs.slide_layouts[6]
    for text in texts:
        slide = prs.slides.add_slide(layout)
        tb = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
        tb.text_frame.text = text
        if with_image:
            img_path = os.path.join(tmp_path, "pic.png")
            if not os.path.isfile(img_path):
                Image.new("RGB", (20, 15), "red").save(img_path)
            slide.shapes.add_picture(img_path, Inches(0.5), Inches(3))
    path = os.path.join(tmp_path, "src.pptx")
    prs.save(path)
    return path


def test_splice_single_slide_appends_and_preserves_text(tmp_path):
    src_path = _make_src_pptx(tmp_path, ["hello from src"])
    dst = Presentation()
    before = len(dst.slides)

    new_slides = splice_external_slides(dst, src_path)

    assert len(new_slides) == 1
    assert len(dst.slides) == before + 1
    assert _slide_texts(dst.slides[-1]) == ["hello from src"]


def test_splice_result_reopens_cleanly(tmp_path):
    src_path = _make_src_pptx(tmp_path, ["page one", "page two"])
    dst = Presentation()
    splice_external_slides(dst, src_path)

    out_path = os.path.join(tmp_path, "out.pptx")
    dst.save(out_path)

    reopened = Presentation(out_path)
    assert len(reopened.slides) == 2
    assert _slide_texts(reopened.slides[0]) == ["page one"]
    assert _slide_texts(reopened.slides[1]) == ["page two"]


def test_splice_specific_indices_only(tmp_path):
    src_path = _make_src_pptx(tmp_path, ["a", "b", "c"])
    dst = Presentation()

    new_slides = splice_external_slides(dst, src_path, slide_indices=[2, 0])

    assert [_slide_texts(s) for s in new_slides] == [["c"], ["a"]]
    assert len(dst.slides) == 2


def test_splice_shares_cloned_layout_across_slides_from_same_source(tmp_path):
    src_path = _make_src_pptx(tmp_path, ["a", "b"])
    dst = Presentation()
    before_masters = len(dst.slide_masters)

    new_slides = splice_external_slides(dst, src_path)

    # 两页共用同一个源 layout/master，克隆时应该只挂一份新 master，不是两份。
    assert len({s.slide_layout.part for s in new_slides}) == 1
    assert len(dst.slide_masters) == before_masters + 1


def test_splice_preserves_embedded_image_bytes(tmp_path):
    src_path = _make_src_pptx(tmp_path, ["with pic"], with_image=True)
    dst = Presentation()

    splice_external_slides(dst, src_path)

    pics = [sh for sh in dst.slides[-1].shapes if sh.shape_type == 13]  # MSO_SHAPE_TYPE.PICTURE
    assert len(pics) == 1
    with open(os.path.join(tmp_path, "pic.png"), "rb") as f:
        original_bytes = f.read()
    assert pics[0].image.blob == original_bytes


def test_splice_skips_customxml_dependency(tmp_path):
    # 回归：真实复现过的"需要修复"——某个 slide/master 关联的 customXml
    # （设计工具自己的版式目录元数据）只克隆到"这一页用到的那部分"，缺了它
    # 所属的完整目录/索引，PowerPoint 打开时会把这批不完整的 customXml 整批
    # 清掉触发修复。这类元数据不影响幻灯片视觉内容，直接不克隆。
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
    from pptx.opc.package import XmlPart
    from pptx.oxml import parse_xml
    from pptx.oxml.ns import qn

    src_path = _make_src_pptx(tmp_path, ["with customxml"])
    src = Presentation(src_path)
    slide_part = src.slides[0].part
    package = slide_part.package
    cx_partname = package.next_partname("/customXml/item%d.xml")
    cx_part = XmlPart(cx_partname, "application/xml", package,
                      parse_xml(b'<root xmlns="http://example.com"/>'))
    rId = slide_part.relate_to(cx_part, RT.CUSTOM_XML)
    # r:id 在 <p:custData> 上是必填属性 —— 真实场景是这个元素本来指着一个
    # customXml 关系；如果只删关系不删元素，会留下一个 schema 非法的空壳。
    cSld = slide_part._element.find(qn("p:cSld"))
    custDataLst = cSld.makeelement(qn("p:custDataLst"), {})
    custData = custDataLst.makeelement(qn("p:custData"), {qn("r:id"): rId})
    custDataLst.append(custData)
    cSld.append(custDataLst)
    src.save(src_path)

    dst = Presentation()
    splice_external_slides(dst, src_path)
    out_path = os.path.join(tmp_path, "out.pptx")
    dst.save(out_path)

    with zipfile.ZipFile(out_path) as zf:
        assert not any("customxml" in n.lower() for n in zf.namelist())
        slide_xml = zf.read("ppt/slides/slide1.xml").decode("utf-8")
        # custData 的 r:id 是必填属性——关系没克隆过来的话，元素本身也要一起
        # 删掉，不能留一个缺 r:id 的空壳（那样才是 schema 非法）。
        assert "custData" not in slide_xml


def test_splice_new_master_id_does_not_collide_with_existing_layout_ids(tmp_path):
    # 真实复现过、且是"需要修复"长期调查的最终根因：新 master 的 id 只在
    # p:sldMasterId 这一份列表里取 max()+1，但 PowerPoint 把 p:sldMasterId
    # 和每个已有 master 自己 p:sldLayoutIdLst 里的 p:sldLayoutId 当成同一个
    # 共享的 32 位 id 空间——两者是各自独立分配的，"下一个"号码很容易撞上
    # 一个完全不相关的、已有 master 的 layout id。默认 python-pptx 模板本身
    # 就是最小复现：sldMasterId=2147483648，它自己的 11 个 layout 的 id 是
    # 2147483649..2147483659，naive max(sldMasterId)+1=2147483649 直接撞上
    # 第一个 layout id。
    from pptx.oxml.ns import qn

    src_path = _make_src_pptx(tmp_path, ["new master"])
    dst = Presentation()

    existing_layout_ids = {
        int(v) for v in
        dst.slide_masters[0]._element.xpath(".//p:sldLayoutIdLst/p:sldLayoutId/@id")
    }
    existing_master_ids = {
        int(el.get("id")) for el in
        dst.part._element.get_or_add_sldMasterIdLst().sldMasterId_lst
    }

    splice_external_slides(dst, src_path)

    all_master_ids = {
        int(el.get("id")) for el in
        dst.part._element.get_or_add_sldMasterIdLst().sldMasterId_lst
    }
    new_master_ids = all_master_ids - existing_master_ids
    assert len(new_master_ids) == 1
    new_id = next(iter(new_master_ids))
    assert new_id not in existing_layout_ids
    assert new_id not in existing_master_ids


def test_splice_adds_raster_fallback_to_bare_svg_blip(tmp_path):
    # 真实源模板里常见一种写法：<a:blip> 外层没有 r:embed/r:link，只在
    # <a:extLst>/asvg:svgBlip 里挂了 SVG。PowerPoint 自己插入 SVG 图片时
    # 一定会同时生成一份 raster fallback 挂在外层 r:embed 上；源文件缺了
    # 这个 fallback 疑似是"需要修复"的诱因之一，拼接时补上。
    import io

    from PIL import Image
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
    from pptx.oxml.ns import qn

    src_path = _make_src_pptx(tmp_path, ["with bare svg blip"], with_image=True)
    src = Presentation(src_path)
    slide_part = src.slides[0].part
    package = slide_part.package

    svg_bytes = (b'<svg xmlns="http://www.w3.org/2000/svg" width="10" '
                 b'height="10"><rect width="10" height="10" fill="red"/></svg>')
    svg_partname = package.next_partname("/ppt/media/image%d.svg")
    from pptx.opc.package import Part
    svg_part = Part(svg_partname, "image/svg+xml", package, svg_bytes)
    svg_rid = slide_part.relate_to(svg_part, RT.IMAGE)

    r_ns_ = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
    blip = slide_part._element.find(f".//{qn('a:blip')}")
    del blip.attrib[r_ns_ + "embed"]  # 模拟"外层没有 r:embed，只靠 svgBlip"的源写法
    ext_lst = blip.makeelement(qn("a:extLst"), {})
    ext = ext_lst.makeelement(qn("a:ext"), {"uri": "{96DAC541-7B7A-43D3-8B79-37D633B846F1}"})
    asvg_ns = "http://schemas.microsoft.com/office/drawing/2016/SVG/main"
    svg_blip = ext.makeelement(f"{{{asvg_ns}}}svgBlip", {qn("r:embed"): svg_rid})
    ext.append(svg_blip)
    ext_lst.append(ext)
    blip.append(ext_lst)
    src.save(src_path)

    dst = Presentation()
    splice_external_slides(dst, src_path)
    out_path = os.path.join(tmp_path, "out.pptx")
    dst.save(out_path)

    new_slide_part = dst.slides[-1].part
    new_pic = new_slide_part._element.find(f".//{qn('p:pic')}")
    new_blip = new_pic.find(f".//{qn('a:blip')}")
    r_ns = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
    embed_rid = new_blip.get(r_ns + "embed")
    assert embed_rid, "outer a:blip should get a resolvable r:embed fallback"

    fallback_part = new_slide_part.rels[embed_rid].target_part
    assert fallback_part.content_type == "image/png"
    with io.BytesIO(fallback_part.blob) as buf:
        Image.open(buf).verify()

    # svgBlip 扩展的 r:embed 也要被正确重映射到新拼接出的 svg 关系上，不是
    # 被外层 fallback 顶掉。
    new_svg_blip = new_blip.find(f".//{{{asvg_ns}}}svgBlip")
    svg_embed_rid = new_svg_blip.get(r_ns + "embed")
    assert new_slide_part.rels[svg_embed_rid].target_part.content_type == "image/svg+xml"


def test_fix_app_properties_slide_count(tmp_path):
    # 回归：docProps/app.xml 的 <Slides>/<Notes> 计数 python-pptx 从不维护——
    # 默认空白模板永远是 0，真实遇到过 PowerPoint 拿这个跟实际页数不一致时
    # 提示"需要修复"（schema 完全合法，validator 查不出来）。
    src_path = _make_src_pptx(tmp_path, ["a", "b", "c"])
    prs = Presentation(src_path)

    fix_app_properties_slide_count(prs)
    out_path = os.path.join(tmp_path, "fixed.pptx")
    prs.save(out_path)

    with zipfile.ZipFile(out_path) as zf:
        xml = zf.read("docProps/app.xml").decode("utf-8")
    assert re.search(r"<Slides>3</Slides>", xml)
    assert re.search(r"<Notes>3</Notes>", xml)


def test_remove_slide_drops_correct_slide(tmp_path):
    src_path = _make_src_pptx(tmp_path, ["keep me", "drop me"])
    prs = Presentation(src_path)

    remove_slide(prs, 1)

    assert len(prs.slides) == 1
    assert _slide_texts(prs.slides[0]) == ["keep me"]


def test_partname_allocator_treats_names_as_case_insensitive():
    # 回归：真实拼接过程中出现过 "customXml/item291.xml" (源) vs
    # "customXML/item291.xml" (目标模板自带) —— OPC part 名唯一性按大小写不
    # 敏感判断，.NET 的 System.IO.Packaging 会当成同一个 part 冲突，PowerPoint
    # 打开时报"需要修复"。之前的实现按大小写敏感字符串比较，会把这判成"没冲突"。
    class _FakePart:
        def __init__(self, name):
            self.partname = name

    class _FakePackage:
        def iter_parts(self):
            return [_FakePart("customXML/item1.xml")]

    alloc = _PartnameAllocator(_FakePackage())
    got = alloc.next("customXml/item%d.xml")

    assert got.lower() != "customxml/item1.xml"
    assert got == "customXml/item2.xml"


def test_remove_slide_result_reopens_cleanly(tmp_path):
    src_path = _make_src_pptx(tmp_path, ["a", "b", "c"])
    prs = Presentation(src_path)
    remove_slide(prs, 0)
    out_path = os.path.join(tmp_path, "out.pptx")
    prs.save(out_path)

    reopened = Presentation(out_path)
    assert len(reopened.slides) == 2
    assert _slide_texts(reopened.slides[0]) == ["b"]
