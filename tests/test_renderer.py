"""renderer：渲染出有效 pptx、无负几何、相册/时间轴回归。"""
from pptx.util import Inches
from pptx import Presentation

from src.models import ChartData, DeckPlan, SlideData, TableData
from src.renderer import render


def _base_plan():
    return DeckPlan(title="测试", theme_name="business", slides=[
        SlideData(kind="title", title="封面"),
        SlideData(kind="bullets", title="要点", bullets=["甲", "乙", "丙"]),
        SlideData(kind="table", title="表格",
                  table=TableData(headers=["h"], rows=[["1"]])),
        SlideData(kind="chart", title="图表",
                  chart=ChartData(chart_type="column", categories=["Q1"],
                                  series=[{"name": "A", "values": [1]}])),
        SlideData(kind="summary", title="总结", bullets=["要点1"]),
        SlideData(kind="thanks"),
    ])


def _assert_no_negative_geometry(prs):
    for slide in prs.slides:
        for sh in slide.shapes:
            assert sh.width is None or sh.width >= 0
            assert sh.height is None or sh.height >= 0


def test_render_base_deck(tmp_path):
    out = str(tmp_path / "deck.pptx")
    render(_base_plan(), out)
    prs = Presentation(out)
    assert len(prs.slides) == 6
    _assert_no_negative_geometry(prs)


def test_render_photo_timeline_no_negative(tmp_path, sample_img_with_exif):
    # 回归：时间轴下方卡片的连线高度曾为负 -> 损坏 pptx
    plan = DeckPlan(title="t", theme_name="education", slides=[
        SlideData(kind="title", title="封面"),
        SlideData(kind="photo_timeline", title="时间轴",
                  images=[sample_img_with_exif] * 4),
        SlideData(kind="thanks"),
    ])
    out = str(tmp_path / "tl.pptx")
    render(plan, out)
    prs = Presentation(out)
    assert len(prs.slides) == 3
    _assert_no_negative_geometry(prs)


def test_render_gallery_with_plain_images(tmp_path, sample_img):
    plan = DeckPlan(title="t", theme_name="business", slides=[
        SlideData(kind="title", title="封面"),
        SlideData(kind="gallery", title="相册", images=[sample_img] * 3),
        SlideData(kind="thanks"),
    ])
    out = str(tmp_path / "g.pptx")
    render(plan, out)
    prs = Presentation(out)
    _assert_no_negative_geometry(prs)


def test_render_chart_has_no_negative_axis_ids(tmp_path):
    # 回归：python-pptx 自己的图表 XML 模板里 axId/crossAx 是负数十进制字面量，
    # 违反 schema 的 UInt32 类型，PowerPoint 打开含图表的 deck 时可能提示"需要
    # 修复"（真实复现过：SRDP 全流程生成的 deck 里第 5 张图表触发过）。
    import re

    from pptx import Presentation as _Presentation

    plan = DeckPlan(title="t", theme_name="business", slides=[
        SlideData(kind="title", title="封面"),
        SlideData(kind="chart", title="图表",
                  chart=ChartData(chart_type="column", categories=["a", "b"],
                                  series=[{"name": "x", "values": [1, 2]}])),
        SlideData(kind="thanks"),
    ])
    out = str(tmp_path / "chart.pptx")
    render(plan, out)

    _Presentation(out)  # 能重新打开
    import zipfile
    with zipfile.ZipFile(out) as zf:
        chart_xmls = [n for n in zf.namelist()
                     if n.startswith("ppt/charts/chart") and n.endswith(".xml")]
        assert chart_xmls
        for name in chart_xmls:
            xml = zf.read(name).decode("utf-8")
            for m in re.finditer(r'(?:axId|crossAx) val="(-?\d+)"', xml):
                n = int(m.group(1))
                assert 0 <= n <= 2147483647, f"{name}: axId/crossAx out of range {n}"


def _make_template(tmp_path, width_in: float, height_in: float) -> str:
    """造一个真实的、真实尺寸不同于渲染器默认 13.333x7.5in 的模板 pptx。"""
    prs = Presentation()
    prs.slide_width = Inches(width_in)
    prs.slide_height = Inches(height_in)
    path = str(tmp_path / "template.pptx")
    prs.save(path)
    return path


def _assert_within_canvas(prs) -> None:
    """只检查我们自己画的形状（跳过继承自 layout 的占位符——renderer 不缩放
    也不保证它们的几何，那是模板设计者的责任；这里用的测试模板是拿默认
    Presentation() 硬改 slide_width/height 造的，layout 占位符并未跟着改，
    不能代表真实模板（如 MainLiveDoc.pptx）的情况）。"""
    W, H = prs.slide_width, prs.slide_height
    for slide in prs.slides:
        for sh in slide.shapes:
            if sh.left is None or sh.width is None or sh.is_placeholder:
                continue
            assert sh.left + sh.width <= W + Inches(0.05)
            assert sh.top + sh.height <= H + Inches(0.05)


def test_render_with_template_uses_its_real_canvas_size(tmp_path):
    template = _make_template(tmp_path, 10.0, 5.625)
    out = str(tmp_path / "templated.pptx")

    render(_base_plan(), out, template_path=template)

    prs = Presentation(out)
    assert prs.slide_width == Inches(10.0)
    assert prs.slide_height == Inches(5.625)
    assert len(prs.slides) == 6
    _assert_no_negative_geometry(prs)
    _assert_within_canvas(prs)


def test_render_with_template_scales_cards_that_overflowed_before_fix(tmp_path):
    # 回归：早期实现把画布切到模板真实尺寸后不缩放形状，3 行卡片会超出更矮的
    # 画布（见开发过程中 avail_h=4.9 硬编码越界的问题）。
    template = _make_template(tmp_path, 10.0, 5.625)
    plan = DeckPlan(title="t", theme_name="business", slides=[
        SlideData(kind="title", title="封面"),
        SlideData(kind="cards", title="卡片",
                  bullets=[f"要点{i}" for i in range(9)]),
        SlideData(kind="thanks"),
    ])
    out = str(tmp_path / "cards.pptx")

    render(plan, out, template_path=template)

    prs = Presentation(out)
    _assert_within_canvas(prs)


def test_render_without_template_keeps_default_canvas_after_a_template_call(tmp_path):
    # 回归：renderer 内部曾用 global 改写 SLIDE_W/H/MARGIN，若不在无模板分支
    # 显式重置，会污染同一进程里后续不带模板的渲染。
    template = _make_template(tmp_path, 10.0, 5.625)
    render(_base_plan(), str(tmp_path / "templated.pptx"), template_path=template)

    out = str(tmp_path / "plain.pptx")
    render(_base_plan(), out)

    prs = Presentation(out)
    assert prs.slide_width == Inches(13.333)
    assert prs.slide_height == Inches(7.5)
