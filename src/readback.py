"""读回内容 QA —— Phase D（对应 ppt master 的 readBackQA）。

渲染完成后重新打开 pptx，把「计划预期」与「实际渲染」逐页核对：

  - 内容页(kind 属于内容型)渲染后却没有任何文字/表格/图表/图片 -> empty_render（空壳页）
  - 预期 table 但实际没有表格形状 / 预期 chart 但实际没有图表      -> missing_object
  - 预期 gallery/image/figure 但实际没有任何图片                  -> missing_image

``auto_fix`` 把空壳页/残缺页从计划中剔除（结构页 title/section/thanks 除外），
返回新计划供重新渲染；空壳页本来也渲染不出有效内容，剔除后整份更干净。

与 src/quality.analyze（几何：溢出/重叠）互补：本模块管「内容完整性」。
"""
from __future__ import annotations

import copy
from collections import Counter
from dataclasses import dataclass

from pptx import Presentation

from .models import DeckPlan

CONTENT_KINDS = {"agenda", "bullets", "cards", "timeline", "table", "chart",
                 "gallery", "image", "figure", "summary", "photo_timeline"}
STRUCT_KINDS = {"title", "section", "thanks"}


@dataclass
class QAIssue:
    slide: int          # 1-based
    kind: str           # empty_render / missing_object / missing_image
    message: str
    severity: str = "error"


def _shape_content(sh) -> set[str]:
    """返回该 shape 的内容标记（text/table/chart/pic）。"""
    tags: set[str] = set()
    if sh.has_text_frame and sh.text_frame.text.strip():
        tags.add("text")
    if sh.has_table:
        tags.add("table")
    if sh.has_chart:
        tags.add("chart")
    if sh.shape_type == 13:  # PICTURE
        tags.add("pic")
    return tags


def compare_plan_render(plan: DeckPlan, path: str) -> list[QAIssue]:
    """计划预期 vs 实际渲染逐页核对，返回问题列表（无 Key 依赖，纯本地）。"""
    prs = Presentation(path)
    issues: list[QAIssue] = []
    for i, sd in enumerate(plan.slides, 1):
        if i > len(prs.slides):
            break
        actual: set[str] = set()
        for sh in prs.slides[i - 1].shapes:
            actual |= _shape_content(sh)
        kind = sd.kind

        if kind in CONTENT_KINDS and not actual:
            issues.append(QAIssue(
                i, "empty_render",
                f"页 {i}「{sd.title}」计划为 {kind}，但渲染后无任何内容"))
            continue
        if kind == "table" and "table" not in actual:
            issues.append(QAIssue(
                i, "missing_object",
                f"页 {i}「{sd.title}」预期表格，实际未渲染出表格"))
        if kind == "chart" and "chart" not in actual:
            issues.append(QAIssue(
                i, "missing_object",
                f"页 {i}「{sd.title}」预期图表，实际未渲染出图表"))
        if kind in ("gallery", "image", "photo_timeline", "figure") and "pic" not in actual:
            issues.append(QAIssue(
                i, "missing_image",
                f"页 {i}「{sd.title}」预期图片，实际无图片"))
        if kind in ("bullets", "cards", "timeline", "summary", "agenda", "figure") \
                and sd.bullets and "text" not in actual:
            issues.append(QAIssue(
                i, "empty_render",
                f"页 {i}「{sd.title}」有 {len(sd.bullets)} 条要点，但渲染后无文字"))
    return issues


def auto_fix(plan: DeckPlan, issues: list[QAIssue]) -> tuple[DeckPlan, int]:
    """从计划中剔除空壳/残缺内容页（结构页保留）；返回 (新 plan, 剔除页数)。"""
    drop = {it.slide for it in issues
            if it.kind in ("empty_render", "missing_object", "missing_image")
            and it.severity == "error"}
    if not drop:
        return plan, 0
    new_plan = copy.deepcopy(plan)
    kept = [sd for i, sd in enumerate(new_plan.slides, 1)
            if not (i in drop and sd.kind not in STRUCT_KINDS)]
    if len(kept) < 2:
        return plan, 0  # 兜底：不冒险删到只剩 1 页
    new_plan.slides = kept
    return new_plan, len(plan.slides) - len(kept)


def summary(issues: list[QAIssue]) -> str:
    if not issues:
        return "通过：所有页面内容与计划一致。"
    cnt = Counter(i.kind for i in issues)
    return "发现 " + ", ".join(f"{k}×{v}" for k, v in cnt.items())
