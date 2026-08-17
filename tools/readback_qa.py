"""PPTX 读回 QA：重新打开生成的 pptx，逐页核对内容与几何（Phase D 雏形）。

输出每页：shape 数 / 文本框数 / 字符数 / 内容类型(chart/table/pic) / 标题，
最后用 src.quality.analyze 做溢出/重叠/文本溢出检查。

用法：
    python tools/readback_qa.py output/demo_final.pptx
"""
from __future__ import annotations

import argparse
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from pptx import Presentation


def _slide_title(s) -> str:
    """取页面上第一个有文字的文本框作为标题（近似）。"""
    for sh in s.shapes:
        if sh.has_text_frame:
            t = sh.text_frame.text.strip()
            if t:
                return t.split("\n")[0][:30]
    return ""


def qa(path: str) -> int:
    prs = Presentation(path)
    print(f"尺寸: {prs.slide_width / 914400:.2f} x {prs.slide_height / 914400:.2f} in"
          f" | 页数: {len(prs.slides)}")
    for i, s in enumerate(prs.slides, 1):
        kinds: list[str] = []
        ntext = nchar = 0
        for sh in s.shapes:
            if sh.has_chart:
                kinds.append("chart")
            elif sh.has_table:
                kinds.append("table")
            elif sh.shape_type == 13:  # PICTURE
                kinds.append("pic")
            if sh.has_text_frame:
                t = sh.text_frame.text.strip()
                if t:
                    ntext += 1
                    nchar += len(t)
        kind = ",".join(dict.fromkeys(kinds)) or "shapes"
        print(f"{i:02d} shapes={len(s.shapes):2d} textbox={ntext:2d} "
              f"chars={nchar:4d} [{kind:10}] {_slide_title(s)}")

    # 几何质检（复用 quality.analyze）
    try:
        from src.quality import analyze, summary
        issues = analyze(path)
        print("\n[几何质检]", summary(issues))
        for it in issues[:10]:
            print(f"  - 页{it.slide} [{it.kind}] {it.detail}")
    except Exception as e:
        print(f"\n[几何质检] 跳过：{e}")
    return 0


def main() -> int:
    ap = argparse.ArgumentParser(description="PPTX 读回 QA")
    ap.add_argument("deck", help="输入 .pptx 路径")
    args = ap.parse_args()
    return qa(args.deck)


if __name__ == "__main__":
    raise SystemExit(main())
